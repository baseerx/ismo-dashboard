"""Leave and attendance reports as a downloadable Excel or PDF file.

Both formats are built from the same rows so a printed PDF and an exported
spreadsheet can never tell different stories. Everything is produced in memory
and streamed: reports contain one employee's HR record, and leaving copies on
the server's disk is a leak waiting to happen.
"""

import io
from datetime import date, datetime
from typing import Any, Dict, List, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ORGANISATION = "Independent System & Market Operator (ISMO)"

# Kept in step with the dashboard's own palette so a report looks like it came
# from the same system.
INK = colors.HexColor("#1C2430")
ACCENT = colors.HexColor("#0F1B38")
ACCENT_SOFT = colors.HexColor("#E8F6F9")
RULE = colors.HexColor("#D5DAE2")

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"

_HEADER_FILL = PatternFill("solid", fgColor="0F1B38")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_TITLE_FONT = Font(bold=True, size=13)
_THIN = Side(style="thin", color="D5DAE2")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# ------------------------------------------------------------------ shaping

LEAVE_COLUMNS = (
    ("leave_type", "Leave Type", 26),
    ("start_date", "From", 14),
    ("end_date", "To", 14),
    ("days", "Days", 8),
    ("status", "Status", 12),
    ("approver_name", "Section Head", 26),
    ("reason", "Reason", 40),
)

ATTENDANCE_COLUMNS = (
    ("date", "Date", 14),
    ("checkin_time", "Check In", 11),
    ("checkout_time", "Check Out", 11),
    ("late_status", "Arrival", 11),
    ("early_status", "Departure", 12),
    ("status", "Day Status", 22),
)

OFFICIAL_WORK_COLUMNS = (
    ("leave_type", "Type", 24),
    ("start_date", "From", 14),
    ("end_date", "To", 14),
    ("status", "Status", 12),
    ("reason", "Purpose", 44),
)

COLUMNS = {
    "leave": LEAVE_COLUMNS,
    "attendance": ATTENDANCE_COLUMNS,
    "official_work": OFFICIAL_WORK_COLUMNS,
}

TITLES = {
    "leave": "Leave Report",
    "attendance": "Attendance Report",
    "official_work": "Official Work Report",
}


def _cell(value: Any) -> str:
    if value is None or value == "":
        return "-"
    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y %H:%M")
    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")
    return str(value)


def _table(subject: str, rows: Sequence[Dict[str, Any]]) -> Tuple[List[str], List[List[str]], List[int]]:
    columns = COLUMNS[subject]
    headers = [heading for _, heading, _ in columns]
    widths = [width for _, _, width in columns]
    body = [[_cell(row.get(key)) for key, _, _ in columns] for row in rows]
    return headers, body, widths


def filename_for(subject: str, employee_name: str, start: date, end: date, extension: str) -> str:
    safe_name = "".join(ch if ch.isalnum() else "-" for ch in employee_name).strip("-")
    return f"{subject}-{safe_name}-{start:%Y%m%d}-{end:%Y%m%d}.{extension}".lower()


# -------------------------------------------------------------------- excel

def build_excel(
    subject: str,
    employee: Dict[str, Any],
    period: Tuple[date, date],
    rows: Sequence[Dict[str, Any]],
    summary: Dict[str, Any] | None = None,
) -> bytes:
    headers, body, widths = _table(subject, rows)
    start, end = period

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = TITLES[subject].replace(" Report", "")[:31]

    sheet["A1"] = f"{ORGANISATION} — {TITLES[subject]}"
    sheet["A1"].font = _TITLE_FONT
    sheet["A2"] = (
        f"{employee.get('name', '-')}  (ERP {employee.get('erp_id', '-')})"
        f"   |   {employee.get('section') or '-'}"
    )
    sheet["A3"] = f"Period: {start:%d %b %Y} to {end:%d %b %Y}"
    sheet["A4"] = f"Generated: {datetime.now():%d %b %Y %H:%M}"

    for row_index in (1, 2, 3, 4):
        sheet.merge_cells(
            start_row=row_index, start_column=1, end_row=row_index, end_column=len(headers)
        )

    header_row = 6
    for column, heading in enumerate(headers, start=1):
        cell = sheet.cell(row=header_row, column=column, value=heading)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _BORDER

    for offset, line in enumerate(body, start=header_row + 1):
        for column, value in enumerate(line, start=1):
            cell = sheet.cell(row=offset, column=column, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=column == len(headers))

    if not body:
        sheet.cell(row=header_row + 1, column=1, value="No records in this period.")

    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    # Freeze the header so a long report stays readable while scrolling.
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    if summary:
        summary_row = header_row + len(body) + 2
        sheet.cell(row=summary_row, column=1, value="Summary").font = Font(bold=True)
        for offset, (label, value) in enumerate(summary.items(), start=summary_row + 1):
            sheet.cell(row=offset, column=1, value=label.replace("_", " ").title())
            sheet.cell(row=offset, column=2, value=value)

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


# ---------------------------------------------------------------------- pdf

def build_pdf(
    subject: str,
    employee: Dict[str, Any],
    period: Tuple[date, date],
    rows: Sequence[Dict[str, Any]],
    summary: Dict[str, Any] | None = None,
) -> bytes:
    headers, body, widths = _table(subject, rows)
    start, end = period

    stream = io.BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=landscape(A4) if len(headers) > 5 else A4,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=f"{TITLES[subject]} — {employee.get('name', '')}",
        author=ORGANISATION,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Title"], fontSize=15, textColor=ACCENT, spaceAfter=2
    )
    meta_style = ParagraphStyle(
        "ReportMeta", parent=styles["Normal"], fontSize=9, textColor=INK, leading=13
    )
    cell_style = ParagraphStyle("Cell", parent=styles["Normal"], fontSize=8, leading=10)

    story: List[Any] = [
        Paragraph(TITLES[subject], title_style),
        Paragraph(ORGANISATION, meta_style),
        Spacer(1, 6),
        Paragraph(
            f"<b>{employee.get('name', '-')}</b> &nbsp;·&nbsp; ERP {employee.get('erp_id', '-')}"
            f" &nbsp;·&nbsp; {employee.get('section') or '-'}"
            f"<br/>Period: <b>{start:%d %b %Y}</b> to <b>{end:%d %b %Y}</b>"
            f"<br/>Generated: {datetime.now():%d %b %Y %H:%M}",
            meta_style,
        ),
        Spacer(1, 10),
    ]

    if body:
        # Paragraphs rather than bare strings so long reasons wrap instead of
        # running off the page.
        data = [headers] + [[Paragraph(value, cell_style) for value in line] for line in body]
        total = sum(widths)
        available = document.width
        table = Table(
            data,
            colWidths=[available * (width / total) for width in widths],
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), ACCENT),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 8.5),
                    ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("GRID", (0, 0), (-1, -1), 0.4, RULE),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ACCENT_SOFT]),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("No records in this period.", meta_style))

    if summary:
        story.append(Spacer(1, 12))
        story.append(Paragraph("<b>Summary</b>", meta_style))
        summary_rows = [
            [label.replace("_", " ").title(), str(value)] for label, value in summary.items()
        ]
        summary_table = Table(summary_rows, colWidths=[60 * mm, 30 * mm])
        summary_table.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.4, RULE),
                    ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                    ("BACKGROUND", (0, 0), (0, -1), ACCENT_SOFT),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(summary_table)

    document.build(story)
    return stream.getvalue()


def build(
    report_format: str,
    subject: str,
    employee: Dict[str, Any],
    period: Tuple[date, date],
    rows: Sequence[Dict[str, Any]],
    summary: Dict[str, Any] | None = None,
) -> Tuple[bytes, str, str]:
    """Returns (file bytes, filename, media type)."""
    if subject not in COLUMNS:
        raise ValueError(f"Unknown report subject: {subject}")

    if report_format == "excel":
        payload = build_excel(subject, employee, period, rows, summary)
        extension, media_type = "xlsx", XLSX_MEDIA_TYPE
    elif report_format == "pdf":
        payload = build_pdf(subject, employee, period, rows, summary)
        extension, media_type = "pdf", PDF_MEDIA_TYPE
    else:
        raise ValueError(f"Unknown report format: {report_format}")

    name = filename_for(subject, str(employee.get("name", "employee")), period[0], period[1], extension)
    return payload, name, media_type
